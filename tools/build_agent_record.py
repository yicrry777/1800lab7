#!/usr/bin/env python3
"""Fill the retained Assignment 7 Agent Record template in place."""

from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "docs" / "Emerging_Technologies_Agent_Record_Template.docx"
OUTPUT = ROOT / "deliverables" / "Shuoying_Wang_Assignment_7_Organizational_Adoption_Agent_Record.docx"


def set_text(cell, text, size=9.5):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def fill_table(table, values, size=9.5):
    assert len(table.rows) - 1 == len(values)
    repeat_header(table.rows[0])
    for row, value in zip(table.rows[1:], values):
        set_text(row.cells[1], value, size=size)
        prevent_split(row)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)

    doc.paragraphs[1].text = "Organizational Adoption Agent Record"
    doc.paragraphs[1].style = doc.styles["Title"]
    doc.paragraphs[2].text = "MASY1-GC 1800 | Assignment 7"
    doc.paragraphs[5].paragraph_format.page_break_before = True
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            if "Specialist Agent Record Template" in paragraph.text:
                paragraph.text = "Organizational Adoption Agent Record"

    fill_table(doc.tables[1], [
        "Shuoying Wang",
        "Assignment 7 - Organizational Adoption Agent, version 1.0",
        "October 26, 2026",
        "https://github.com/yicrry777/1800lab7",
        "main",
        "a07466a - feat: build and test organizational adoption agent",
        "MASY1800_ET_Agent_Scaffold_v1_0",
    ])

    fill_table(doc.tables[2], [
        "Even if adoption is strategically appropriate, can this organization absorb, govern, and sustain the proposed application?",
        "Multimodal generative AI for short-form social video creation and content governance.",
        "Generate draft 15-60 second TikTok and Instagram marketing videos from approved assets, with human approval and provenance/disclosure controls before publication.",
        "Northstar Consumer Brands, a hypothetical multinational consumer-products company operating in US and EU digital advertising.",
        "Moderate fast-follower: willing to pilot proven enterprise tools, but unwilling to publish unsupervised AI-generated video.",
        "Public synthetic media creates rights, consent, deceptive-content, privacy, platform-policy, and brand risks. Takedown is possible, but spread can precede correction.",
        "Decide now whether to authorize a six-month controlled pilot; close readiness gates within 90 days and make a scale decision at month six.",
    ], size=9.0)

    fill_table(doc.tables[3], [
        "The technology can create useful drafts, but responsible adoption requires governed inputs, competent human evaluation, rights/vendor controls, disclosure or provenance, monitoring, and named accountability.",
        "The application is plausible for bounded draft generation and localization, not autonomous publication. It requires approved assets, rights checks, independent approval, disclosure, provenance testing, and rapid correction.",
        "Northstar is conditionally ready for a pre-publication pilot, not external publication. Ratings: leadership 3; skills 2; workflow 2; infrastructure 1; resources 2; governance 2; culture 0; operating ownership 1.",
        "Authorize only a bounded pre-publication pilot. Block external publication until vendor review, independent approval, end-to-end provenance/disclosure testing, training, incident ownership, and measurement gates are evidenced.",
        "Moderate confidence in the conditional-pilot judgment; low confidence in publication readiness until vendor documents, workflow tests, owner assignments, training results, and workload/culture evidence exist.",
        "A passing end-to-end publication-control test, named owners, approved vendor evidence, trained reviewers, and sustainable workload would support publication. A critical incident or failed control would require deferral.",
    ])

    fill_table(doc.tables[4], [
        "NIST AI 600-1 (2024-07-26); TikTok AI-generated content policy (accessed 2026-10-26); C2PA Specifications (accessed 2026-10-26); Adobe Firefly Video/Content Credentials announcement (2025-02-12); supplied hypothetical case facts.",
        "All scripts ran. Common schema validation passed. Northstar was conditionally ready for a bounded pre-publication pilot, with external publication blocked pending P0 gates.",
        "The general technology finding and application control demands remained identical: governed inputs, human evaluation, rights/vendor controls, disclosure/provenance, monitoring, and accountability.",
        "Readiness changed from conditionally ready at Northstar to not ready for Harbor Glow's 30-day public launch. Actions, owners, timing, and permitted next step changed with capability.",
        "The v0.1 response passed schema validation but produced a generic checklist, did not assess all eight dimensions, and did not bind priorities to observed gaps, owners, or evidence.",
        "Version 1.0 added eight dimension ratings, non-averaging gates, action-quality requirements, exact evidence requests, and an indeterminate outcome when gating information is missing.",
        "Hypothetical cases do not certify a real organization. The agent abstains when vendor terms, provenance tests, workflow evidence, workload/culture data, budget capacity, or operating owners are missing.",
    ])

    fill_table(doc.tables[5], [
        "Yes, for team consideration. It preserves the common contract, responds to materially different organizations, asks for missing evidence, cites authoritative current sources, and passes schema, context-sensitivity, and frozen-core tests. It remains decision support, not legal or security authority.",
        "Carry forward the eight-dimension intake and gate rule. Do not average away a critical publication control. The final orchestrator should capture named owners, independent approval, vendor evidence, provenance/disclosure test results, workload, budget, training, and incident authority.",
        "ChatGPT/Codex interpreted the assignment, drafted and challenged the specialist logic, generated controlled contrast cases, produced structured runs, and automated validation. Consequential guidance was checked against original NIST, TikTok, C2PA, and Adobe sources. Every prompt build, response validation, package test, frozen-core check, and document render was run; the preserved v0.1 failure shows the revision path.",
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
